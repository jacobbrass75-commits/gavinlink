#!/usr/bin/env python3
"""
Google via Safari — slow and steady to avoid CAPTCHA.
20 second delays between searches. One tab. No CAPTCHA if we're patient.
"""
import json, re, time, csv, subprocess, os
import requests, urllib.parse

PROGRESS_PATH = "/Users/josephsullivan/Downloads/llc_manager_progress.json"
CSV_PATH = "/Users/josephsullivan/Downloads/5+ unit farm.csv"
OUTPUT_PATH = "/Users/josephsullivan/Downloads/LLC_Managers_Contact_Info.xlsx"
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
DELAY = 20  # seconds between Google searches


def extract_emails(text):
    found = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    bad = ['google','example','noreply','sentry','email.com','domain','w3.org','schema.org',
           'googleapis','gstatic','microsoft','bing','outlook.com','hotmail','yahoo',
           'aol','mail.com','protonmail','icloud','live.com','sampleemail',
           'facebook.com','twitter.com','instagram.com','yelp.com','.gov']
    return [e for e in set(found) if not any(x in e.lower() for x in bad) and 5 < len(e) < 60]


def extract_phones(text):
    phones = re.findall(r'(?:\+?1[-.\s]?)?\(?(\d{3})\)?[-.\s]?(\d{3})[-.\s]?(\d{4})', text)
    results = []
    for p in phones:
        num = ''.join(p)
        if num != '0000000000' and len(set(num)) > 2:
            results.append(num)
    return list(set(results))


def sjs(js_code):
    with open('/tmp/_sjs_g.js', 'w') as f:
        f.write(js_code)
    r = subprocess.run(['osascript'], input='''
set f to POSIX file "/tmp/_sjs_g.js"
set c to read f as «class utf8»
tell application "Safari" to tell document 1 to do JavaScript c
''', capture_output=True, text=True, timeout=60)
    return r.stdout.strip()


def osa(script):
    r = subprocess.run(['osascript'], input=script, capture_output=True, text=True, timeout=60)
    return r.stdout.strip()


def google_search(name, llc_name='', city=''):
    contact = {'email': None, 'phone': None, 'linkedin': None}
    query = f'"{name}" "{llc_name}" email phone California'
    encoded_q = urllib.parse.quote(query)

    osa(f'tell application "Safari" to set URL of document 1 to "https://www.google.com/search?q={encoded_q}"')
    time.sleep(6)

    page_text = sjs("document.body.innerText;") or ''

    if 'unusual traffic' in page_text.lower() or 'captcha' in page_text.lower():
        print("[CAPTCHA - waiting 120s]", end=' ', flush=True)
        time.sleep(120)
        osa(f'tell application "Safari" to set URL of document 1 to "https://www.google.com/search?q={encoded_q}"')
        time.sleep(6)
        page_text = sjs("document.body.innerText;") or ''
        if 'unusual traffic' in page_text.lower():
            return contact

    emails = extract_emails(page_text)
    phones = extract_phones(page_text)
    if emails:
        contact['email'] = emails[0]
    if phones:
        contact['phone'] = phones[0]

    # LinkedIn
    links = sjs("Array.from(document.querySelectorAll('a')).map(function(a){return a.href}).filter(function(h){return h.indexOf('linkedin.com/in/')>=0}).join('|||');")
    if links:
        for link in links.split('|||'):
            if 'linkedin.com/in/' in link:
                contact['linkedin'] = link
                break

    # Click into results if nothing
    if not contact['email'] and not contact['phone']:
        result_urls = sjs("""
            var links = document.querySelectorAll('div.g a');
            var urls = [];
            links.forEach(function(a) {
                var h = a.href;
                if (h && !h.includes('google.com') && !h.includes('youtube.com') && !h.includes('wikipedia.org') && !h.includes('facebook.com')) {
                    urls.push(h);
                }
            });
            urls.slice(0, 5).join('|||');
        """)
        if result_urls:
            for rurl in result_urls.split('|||')[:3]:
                if not rurl:
                    continue
                try:
                    page = requests.get(rurl, headers={'User-Agent': UA}, timeout=8)
                    if page.status_code == 200:
                        pe = extract_emails(page.text)
                        pp = extract_phones(page.text)
                        if pe and not contact['email']:
                            contact['email'] = pe[0]
                        if pp and not contact['phone']:
                            contact['phone'] = pp[0]
                        if contact['email'] or contact['phone']:
                            break
                except:
                    pass
                time.sleep(0.5)

    return contact


def generate_excel(progress):
    import openpyxl
    from openpyxl.styles import Font
    with open(CSV_PATH, 'r') as f:
        rows = list(csv.DictReader(f))
    llc_data = {}
    for r in rows:
        owner = r['Owner Name'].strip()
        if 'llc' not in owner.lower():
            continue
        if owner not in llc_data:
            llc_data[owner] = {
                'properties': [],
                'mail_address': r.get('Mail Address', '').strip(),
                'mail_city': r.get('Mail Address City', '').strip(),
            }
        addr = f"{r.get('Site Address','').strip()}, {r.get('Site Address City','').strip()}"
        units = r.get('Number of Units', '')
        prop = f"{addr} ({units} units)" if units else addr
        if prop not in llc_data[owner]['properties']:
            llc_data[owner]['properties'].append(prop)
    wb = openpyxl.Workbook()
    headers = ['LLC Name', 'Properties', '# Properties', 'Mail Address', 'Mail City',
               'Manager Name', 'Email', 'Phone', 'LinkedIn', 'Agent Name', 'CEO']
    ws = wb.active
    ws.title = "With Contact Info"
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True)
    cr = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in [m for m in p.get('managers', []) if m.get('email') or m.get('phone')]:
            ws.append([llc, "; ".join(d['properties']), len(d['properties']),
                       d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                       m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            cr += 1
    ws2 = wb.create_sheet("All Managers")
    ws2.append(headers)
    for c in ws2[1]: c.font = Font(bold=True)
    ar = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in p.get('managers', []):
            ws2.append([llc, "; ".join(d['properties']), len(d['properties']),
                        d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                        m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            ar += 1
    for s in [ws, ws2]:
        for col in s.columns:
            ml = max(len(str(c.value or '')[:50]) for c in col)
            s.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
    wb.save(OUTPUT_PATH)
    print(f"\nExcel saved: {OUTPUT_PATH}")
    print(f"Sheet 1 (with contact): {cr} rows | Sheet 2 (all managers): {ar} rows")


def main():
    with open(PROGRESS_PATH, 'r') as f:
        progress = json.load(f)

    cities = {}
    with open(CSV_PATH, 'r') as f:
        for r in csv.DictReader(f):
            owner = r['Owner Name'].strip()
            if owner not in cities:
                cities[owner] = r.get('Mail Address City', '').strip()

    to_process = []
    for llc_name, info in progress.items():
        for i, mgr in enumerate(info.get('managers', [])):
            if not mgr.get('email') and not mgr.get('phone') and len(mgr.get('name', '').split()) >= 2:
                to_process.append((llc_name, i, mgr['name']))

    total = len(to_process)
    print(f"Remaining managers: {total}")
    print(f"Using {DELAY}s delays to avoid CAPTCHA\n")

    # Open Safari
    osa('''
tell application "Safari"
    activate
    if (count of documents) = 0 then
        make new document with properties {URL:"https://www.google.com"}
    end if
end tell
    ''')
    time.sleep(3)

    found = 0
    for idx, (llc_name, mgr_idx, name) in enumerate(to_process):
        print(f"[{idx+1}/{total}] {name}...", end=' ', flush=True)

        contact = google_search(name, llc_name, cities.get(llc_name, ''))
        if contact.get('email') or contact.get('phone'):
            progress[llc_name]['managers'][mgr_idx]['email'] = contact.get('email')
            progress[llc_name]['managers'][mgr_idx]['phone'] = contact.get('phone')
            progress[llc_name]['managers'][mgr_idx]['linkedin'] = contact.get('linkedin')
            found += 1
            print(f"email={contact.get('email')} phone={contact.get('phone')}")
        else:
            print("nothing")

        if (idx + 1) % 20 == 0:
            with open(PROGRESS_PATH, 'w') as f:
                json.dump(progress, f, indent=2)
            print(f"  --- saved | {found}/{idx+1} found ---")

        time.sleep(DELAY)

    with open(PROGRESS_PATH, 'w') as f:
        json.dump(progress, f, indent=2)
    print(f"\nDone: {found}/{total} found via Google")

    print("\nGenerating Excel...")
    generate_excel(progress)


if __name__ == '__main__':
    main()
