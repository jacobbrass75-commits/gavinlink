#!/usr/bin/env python3
"""
Multi-engine contact lookup: Bing + Yahoo + direct page scraping.
No Safari needed — pure HTTP with proper filtering.
"""
import json, re, time, sys, csv
import requests, urllib.parse
from bs4 import BeautifulSoup
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading

PROGRESS_PATH = "/Users/josephsullivan/Downloads/llc_manager_progress.json"
CSV_PATH = "/Users/josephsullivan/Downloads/5+ unit farm.csv"
OUTPUT_PATH = "/Users/josephsullivan/Downloads/LLC_Managers_Contact_Info.xlsx"
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.2.1 Safari/605.1.15'

progress_lock = threading.Lock()


def extract_emails(text):
    found = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    bad = ['google','example','noreply','sentry','email.com','domain','w3.org','schema.org',
           'googleapis','gstatic','microsoft','bing','outlook.com','hotmail','yahoo.com',
           'aol.com','mail.com','protonmail','icloud','me.com','live.com','sampleemail',
           'your@','user@','info@example','test@','name@','address@','e-mail@',
           'feedback@','support@','admin@','webmaster@','contact@bing','privacy@',
           'abuse@','spam@','mailer-daemon','no-reply','donotreply','noreply',
           '.gov','yelp.com','facebook.com','twitter.com','instagram.com']
    result = []
    for e in set(found):
        low = e.lower()
        if any(x in low for x in bad):
            continue
        if len(e) < 6 or len(e) > 60:
            continue
        # Must have reasonable TLD
        parts = e.split('.')
        if len(parts[-1]) > 5:
            continue
        result.append(e)
    return result


def extract_phones(text):
    phones = re.findall(r'(?:\+?1[-.\s]?)?\(?(\d{3})\)?[-.\s]?(\d{3})[-.\s]?(\d{4})', text)
    results = []
    for p in phones:
        num = ''.join(p)
        if num == '0000000000' or len(set(num)) <= 2:
            continue
        # Filter known junk patterns
        if num in ('8007teletext', '8005551212'):
            continue
        results.append(num)
    return list(set(results))


def scrape_page(url, timeout=8):
    """Scrape a page for emails and phones."""
    try:
        resp = requests.get(url, headers={'User-Agent': UA}, timeout=timeout,
                          allow_redirects=True)
        if resp.status_code != 200:
            return [], [], None
        text = resp.text
        emails = extract_emails(text)
        phones = extract_phones(text)
        # LinkedIn
        linkedin = None
        li_matches = re.findall(r'https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+', text)
        if li_matches:
            linkedin = li_matches[0]
        return emails, phones, linkedin
    except:
        return [], [], None


def bing_search(name, llc_name='', city=''):
    contact = {'email': None, 'phone': None, 'linkedin': None}

    queries = [
        f'"{name}" email phone contact',
        f'"{name}" "{llc_name}" contact California',
        f'"{name}" {city or "California"} property manager email',
    ]

    seen_urls = set()

    for query in queries:
        try:
            resp = requests.get(
                f"https://www.bing.com/search?q={urllib.parse.quote(query)}&count=20",
                headers={'User-Agent': UA}, timeout=10
            )
            if resp.status_code != 200:
                continue

            soup = BeautifulSoup(resp.text, 'html.parser')

            # Extract visible text (not HTML template)
            # Target the actual search result snippets
            result_texts = []
            for div in soup.find_all(['li', 'div'], class_=re.compile(r'b_algo|b_ans')):
                result_texts.append(div.get_text(' ', strip=True))

            combined_text = ' '.join(result_texts)

            emails = extract_emails(combined_text)
            phones = extract_phones(combined_text)

            if emails and not contact['email']:
                contact['email'] = emails[0]
            if phones and not contact['phone']:
                contact['phone'] = phones[0]

            # LinkedIn from result URLs
            for a in soup.find_all('a', href=True):
                href = a['href']
                if 'linkedin.com/in/' in href and not contact['linkedin']:
                    clean = re.match(r'(https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+)', href)
                    if clean:
                        contact['linkedin'] = clean.group(1)

            # If nothing yet, try clicking into result pages
            if not contact['email'] and not contact['phone']:
                for a in soup.find_all('a', href=True):
                    href = a['href']
                    if href.startswith('http') and 'bing.com' not in href and 'microsoft.com' not in href:
                        if href in seen_urls:
                            continue
                        seen_urls.add(href)
                        if any(x in href.lower() for x in ['linkedin.com', 'facebook.com', 'youtube.com', 'wikipedia.org', 'yelp.com', 'twitter.com']):
                            continue
                        pe, pp, pli = scrape_page(href)
                        if pe and not contact['email']:
                            contact['email'] = pe[0]
                        if pp and not contact['phone']:
                            contact['phone'] = pp[0]
                        if pli and not contact['linkedin']:
                            contact['linkedin'] = pli
                        if contact['email'] or contact['phone']:
                            break
                        time.sleep(0.3)
        except:
            pass

        if contact['email'] or contact['phone']:
            return contact
        time.sleep(1)

    return contact


def yahoo_search(name, llc_name='', city=''):
    """Search Yahoo as backup engine."""
    contact = {'email': None, 'phone': None, 'linkedin': None}

    query = f'"{name}" "{llc_name}" email phone contact {city or "California"}'
    try:
        resp = requests.get(
            f"https://search.yahoo.com/search?p={urllib.parse.quote(query)}",
            headers={'User-Agent': UA}, timeout=10
        )
        if resp.status_code != 200:
            return contact

        soup = BeautifulSoup(resp.text, 'html.parser')

        # Get result text
        result_texts = []
        for div in soup.find_all(['div', 'li'], class_=re.compile(r'algo|result|compText')):
            result_texts.append(div.get_text(' ', strip=True))

        combined = ' '.join(result_texts)
        emails = extract_emails(combined)
        phones = extract_phones(combined)

        if emails:
            contact['email'] = emails[0]
        if phones:
            contact['phone'] = phones[0]

        # LinkedIn
        for a in soup.find_all('a', href=True):
            href = a['href']
            if 'linkedin.com/in/' in href and not contact['linkedin']:
                clean = re.match(r'(https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+)', href)
                if clean:
                    contact['linkedin'] = clean.group(1)

        # Scrape top results if nothing
        if not contact['email'] and not contact['phone']:
            for a in soup.find_all('a', href=True):
                href = a['href']
                if href.startswith('http') and 'yahoo.com' not in href:
                    if any(x in href.lower() for x in ['linkedin.com', 'facebook.com', 'youtube.com', 'wikipedia.org', 'yelp.com']):
                        continue
                    pe, pp, pli = scrape_page(href)
                    if pe and not contact['email']:
                        contact['email'] = pe[0]
                    if pp and not contact['phone']:
                        contact['phone'] = pp[0]
                    if contact['email'] or contact['phone']:
                        break
                    time.sleep(0.3)
    except:
        pass

    return contact


def search_worker(item, cities):
    llc_name, mgr_idx, name = item
    city = cities.get(llc_name, '')

    # Try Bing first
    contact = bing_search(name, llc_name, city)
    if contact.get('email') or contact.get('phone'):
        return (llc_name, mgr_idx, name, contact, 'BING')

    time.sleep(0.5)

    # Try Yahoo
    contact = yahoo_search(name, llc_name, city)
    if contact.get('email') or contact.get('phone'):
        return (llc_name, mgr_idx, name, contact, 'YAHOO')

    return (llc_name, mgr_idx, name, None, None)


def generate_excel(progress):
    import openpyxl
    from openpyxl.styles import Font

    with open(CSV_PATH, 'r') as f:
        rows = list(csv.DictReader(f))
    llc_data = {}
    for r in rows:
        owner = r['Owner Name'].strip()
        if 'llc' not in owner.lower():
            continue
        if owner not in llc_data:
            llc_data[owner] = {
                'properties': [],
                'mail_address': r.get('Mail Address', '').strip(),
                'mail_city': r.get('Mail Address City', '').strip(),
            }
        addr = f"{r.get('Site Address','').strip()}, {r.get('Site Address City','').strip()}"
        units = r.get('Number of Units', '')
        prop = f"{addr} ({units} units)" if units else addr
        if prop not in llc_data[owner]['properties']:
            llc_data[owner]['properties'].append(prop)

    wb = openpyxl.Workbook()
    headers = ['LLC Name', 'Properties', '# Properties', 'Mail Address', 'Mail City',
               'Manager Name', 'Email', 'Phone', 'LinkedIn', 'Agent Name', 'CEO']

    ws = wb.active
    ws.title = "With Contact Info"
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True)
    cr = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in [m for m in p.get('managers', []) if m.get('email') or m.get('phone')]:
            ws.append([llc, "; ".join(d['properties']), len(d['properties']),
                       d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                       m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            cr += 1

    ws2 = wb.create_sheet("All Managers")
    ws2.append(headers)
    for c in ws2[1]: c.font = Font(bold=True)
    ar = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in p.get('managers', []):
            ws2.append([llc, "; ".join(d['properties']), len(d['properties']),
                        d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                        m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            ar += 1

    for s in [ws, ws2]:
        for col in s.columns:
            ml = max(len(str(c.value or '')[:50]) for c in col)
            s.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
    wb.save(OUTPUT_PATH)
    print(f"\nExcel saved: {OUTPUT_PATH}")
    print(f"Sheet 1 (with contact): {cr} rows | Sheet 2 (all managers): {ar} rows")


def main():
    with open(PROGRESS_PATH) as f:
        progress = json.load(f)

    cities = {}
    with open(CSV_PATH, 'r') as f:
        for r in csv.DictReader(f):
            owner = r['Owner Name'].strip()
            if owner not in cities:
                cities[owner] = r.get('Mail Address City', '').strip()

    to_process = []
    for llc_name, info in progress.items():
        for i, mgr in enumerate(info.get('managers', [])):
            if not mgr.get('email') and not mgr.get('phone') and len(mgr.get('name', '').split()) >= 2:
                to_process.append((llc_name, i, mgr['name']))

    total = len(to_process)
    print(f"Managers needing contacts: {total}")
    print(f"Searching Bing + Yahoo with 4 threads...\n")

    found = 0
    done = 0

    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(search_worker, item, cities): item for item in to_process}
        for future in as_completed(futures):
            llc_name, mgr_idx, name, contact, source = future.result()
            done += 1

            if contact:
                with progress_lock:
                    progress[llc_name]['managers'][mgr_idx]['email'] = contact.get('email')
                    progress[llc_name]['managers'][mgr_idx]['phone'] = contact.get('phone')
                    progress[llc_name]['managers'][mgr_idx]['linkedin'] = contact.get('linkedin')
                found += 1
                print(f"[{done}/{total}] {name} -> {source}: email={contact.get('email')} phone={contact.get('phone')}")
            elif done % 25 == 0:
                print(f"[{done}/{total}] ... ({found} found)")

            if done % 50 == 0:
                with progress_lock:
                    with open(PROGRESS_PATH, 'w') as f:
                        json.dump(progress, f, indent=2)

    with open(PROGRESS_PATH, 'w') as f:
        json.dump(progress, f, indent=2)

    # Final stats
    total_mgrs = 0
    with_contact = 0
    for info in progress.values():
        for m in info.get('managers', []):
            total_mgrs += 1
            if m.get('email') or m.get('phone'):
                with_contact += 1

    print(f"\nSearch done: found {found}/{total} via Bing/Yahoo")
    print(f"Overall: {with_contact}/{total_mgrs} managers have contacts ({100*with_contact/total_mgrs:.1f}%)")

    print("\nGenerating Excel...")
    generate_excel(progress)


if __name__ == '__main__':
    main()
