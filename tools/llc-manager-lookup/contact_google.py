#!/usr/bin/env python3
"""
Contact lookup via Google in Safari.
RocketReach is rate-limited and DuckDuckGo is blocking us,
so Google via Safari is the only working method right now.
Uses 2 Safari tabs in parallel.
"""

import json, re, time, os, csv, subprocess
import requests, urllib.parse
import threading

PROGRESS_PATH = "/Users/josephsullivan/Downloads/llc_manager_progress.json"
CSV_PATH = "/Users/josephsullivan/Downloads/5+ unit farm.csv"
OUTPUT_PATH = "/Users/josephsullivan/Downloads/LLC_Managers_Contact_Info.xlsx"
HEADERS = {'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}

progress_lock = threading.Lock()
print_lock = threading.Lock()


def extract_emails(text):
    found = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    return [e for e in set(found) if not any(x in e.lower() for x in ['google','example','noreply','sentry','email.com','domain','w3.org','schema.org','googleapis','gstatic'])]

def extract_phones(text):
    phones = re.findall(r'(?:\+?1[-.\s]?)?\(?(\d{3})\)?[-.\s]?(\d{3})[-.\s]?(\d{4})', text)
    return list(set(''.join(p) for p in phones))


def sjs(js_code, tab_idx=1):
    tmp = f'/tmp/_sjs_g{tab_idx}.js'
    with open(tmp, 'w') as f:
        f.write(js_code)
    r = subprocess.run(['osascript'], input=f'''
set f to POSIX file "{tmp}"
set c to read f as «class utf8»
tell application "Safari" to tell tab {tab_idx} of window 1 to do JavaScript c
''', capture_output=True, text=True, timeout=60)
    return r.stdout.strip()

def osa(script):
    r = subprocess.run(['osascript'], input=script, capture_output=True, text=True, timeout=60)
    return r.stdout.strip()


def google_search(name, llc_name, city, tab_idx=1):
    """Search Google via Safari for contact info."""
    contact = {'email': None, 'phone': None, 'linkedin': None}

    queries = [
        f'"{name}" "{llc_name}" email phone California',
        f'"{name}" {city or "California"} property manager contact email phone',
    ]

    for query in queries:
        try:
            encoded_q = urllib.parse.quote(query)
            osa(f'tell application "Safari" to set URL of tab {tab_idx} of window 1 to "https://www.google.com/search?q={encoded_q}"')
            time.sleep(4)

            page_text = sjs("document.body.innerText;", tab_idx) or ''

            if 'unusual traffic' in page_text.lower() or 'captcha' in page_text.lower():
                with print_lock:
                    print(f"[T{tab_idx}] CAPTCHA detected, waiting 60s...", flush=True)
                time.sleep(60)
                continue

            emails = extract_emails(page_text)
            phones = extract_phones(page_text)
            if emails and not contact['email']:
                contact['email'] = emails[0]
            if phones and not contact['phone']:
                contact['phone'] = phones[0]

            # LinkedIn
            links = sjs("Array.from(document.querySelectorAll('a')).map(function(a){return a.href}).filter(function(h){return h.indexOf('linkedin.com/in/')>=0}).join('|||');", tab_idx)
            if links:
                for link in links.split('|||'):
                    if 'linkedin.com/in/' in link:
                        contact['linkedin'] = link
                        break

            # Click into top results if nothing found
            if not contact['email'] and not contact['phone']:
                result_urls = sjs("""
                    var links = document.querySelectorAll('div.g a');
                    var urls = [];
                    links.forEach(function(a) {
                        var h = a.href;
                        if (h && !h.includes('google.com') && !h.includes('youtube.com') && !h.includes('wikipedia.org')) {
                            urls.push(h);
                        }
                    });
                    urls.slice(0, 5).join('|||');
                """, tab_idx)
                if result_urls:
                    for rurl in result_urls.split('|||')[:3]:
                        if not rurl:
                            continue
                        try:
                            page = requests.get(rurl, headers=HEADERS, timeout=8)
                            if page.status_code == 200:
                                emails = extract_emails(page.text)
                                phones = extract_phones(page.text)
                                if emails and not contact['email']:
                                    contact['email'] = emails[0]
                                if phones and not contact['phone']:
                                    contact['phone'] = phones[0]
                                if contact['email'] or contact['phone']:
                                    break
                        except:
                            pass
                        time.sleep(0.5)

        except:
            pass

        if contact['email'] or contact['phone']:
            return contact
        time.sleep(2)

    return contact


def worker(tab_idx, work_queue, progress, cities, counters, total):
    """Worker for one Safari tab."""
    while True:
        with progress_lock:
            if not work_queue:
                return
            llc_name, mgr_idx, name = work_queue.pop(0)
            idx = total - len(work_queue)

        city = cities.get(llc_name, '')

        with print_lock:
            print(f"[T{tab_idx}][{idx}/{total}] {name} ({llc_name})...", end=' ', flush=True)

        contact = google_search(name, llc_name, city, tab_idx)

        if contact.get('email') or contact.get('phone'):
            with progress_lock:
                progress[llc_name]['managers'][mgr_idx]['email'] = contact.get('email')
                progress[llc_name]['managers'][mgr_idx]['phone'] = contact.get('phone')
                progress[llc_name]['managers'][mgr_idx]['linkedin'] = contact.get('linkedin')
                counters['found'] += 1
            with print_lock:
                print(f"email={contact.get('email')} phone={contact.get('phone')}")
        else:
            with print_lock:
                print("nothing")

        with progress_lock:
            counters['done'] += 1
            if counters['done'] % 25 == 0:
                with open(PROGRESS_PATH, 'w') as f:
                    json.dump(progress, f, indent=2)
                with print_lock:
                    print(f"\n  --- {counters['done']}/{total} | {counters['found']} contacts found ---\n")


def generate_excel(progress):
    import openpyxl
    from openpyxl.styles import Font

    with open(CSV_PATH, 'r') as f:
        rows = list(csv.DictReader(f))
    llc_data = {}
    for r in rows:
        owner = r['Owner Name'].strip()
        if 'llc' not in owner.lower():
            continue
        if owner not in llc_data:
            llc_data[owner] = {
                'properties': [],
                'mail_address': r.get('Mail Address', '').strip(),
                'mail_city': r.get('Mail Address City', '').strip(),
            }
        addr = f"{r.get('Site Address','').strip()}, {r.get('Site Address City','').strip()}"
        units = r.get('Number of Units', '')
        prop = f"{addr} ({units} units)" if units else addr
        if prop not in llc_data[owner]['properties']:
            llc_data[owner]['properties'].append(prop)

    wb = openpyxl.Workbook()
    headers = ['LLC Name', 'Properties', '# Properties', 'Mail Address', 'Mail City',
               'Manager Name', 'Email', 'Phone', 'LinkedIn', 'Agent Name', 'CEO']

    ws = wb.active
    ws.title = "With Contact Info"
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True)
    cr = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in [m for m in p.get('managers', []) if m.get('email') or m.get('phone')]:
            ws.append([llc, "; ".join(d['properties']), len(d['properties']),
                       d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                       m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            cr += 1

    ws2 = wb.create_sheet("All Managers")
    ws2.append(headers)
    for c in ws2[1]: c.font = Font(bold=True)
    ar = 0
    for llc, d in llc_data.items():
        p = progress.get(llc, {})
        for m in p.get('managers', []):
            ws2.append([llc, "; ".join(d['properties']), len(d['properties']),
                        d['mail_address'], d['mail_city'], m.get('name',''), m.get('email',''),
                        m.get('phone',''), m.get('linkedin',''), p.get('agent',''), p.get('ceo','')])
            ar += 1

    for s in [ws, ws2]:
        for col in s.columns:
            ml = max(len(str(c.value or '')[:50]) for c in col)
            s.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
    wb.save(OUTPUT_PATH)
    print(f"\nExcel saved: {OUTPUT_PATH}")
    print(f"Sheet 1 (with contact): {cr} rows | Sheet 2 (all managers): {ar} rows")


def main():
    with open(PROGRESS_PATH, 'r') as f:
        progress = json.load(f)

    cities = {}
    with open(CSV_PATH, 'r') as f:
        for r in csv.DictReader(f):
            owner = r['Owner Name'].strip()
            if owner not in cities:
                cities[owner] = r.get('Mail Address City', '').strip()

    # Find managers needing contact info
    work_queue = []
    for llc_name, info in progress.items():
        for i, mgr in enumerate(info.get('managers', [])):
            if not mgr.get('email') and not mgr.get('phone') and len(mgr.get('name', '').split()) >= 2:
                work_queue.append((llc_name, i, mgr['name']))

    total = len(work_queue)
    print(f"Managers needing contact info: {total}")
    print(f"Using Google via Safari (2 tabs)\n")

    # Setup Safari with 2 tabs
    osa('''
tell application "Safari"
    activate
    if (count of windows) = 0 then
        make new document with properties {URL:"https://www.google.com"}
    end if
    tell window 1
        if (count of tabs) < 2 then
            make new tab with properties {URL:"https://www.google.com"}
        end if
        set URL of tab 1 to "https://www.google.com"
        set URL of tab 2 to "https://www.google.com"
    end tell
end tell
    ''')
    time.sleep(4)

    counters = {'found': 0, 'done': 0}

    t1 = threading.Thread(target=worker, args=(1, work_queue, progress, cities, counters, total))
    t2 = threading.Thread(target=worker, args=(2, work_queue, progress, cities, counters, total))
    t1.start()
    time.sleep(2)
    t2.start()
    t1.join()
    t2.join()

    with open(PROGRESS_PATH, 'w') as f:
        json.dump(progress, f, indent=2)

    print(f"\n=== CONTACT LOOKUP COMPLETE ===")
    print(f"Found: {counters['found']}/{total}")
    print("\nGenerating Excel...")
    generate_excel(progress)


if __name__ == '__main__':
    main()
