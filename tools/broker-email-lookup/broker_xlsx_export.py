#!/usr/bin/env python3
"""
Build an xlsx matching the original CoStar export but with broker email columns added.
"""
import json, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from paths import INPUT_CSV, OUTPUT_XLSX, PROGRESS_PATH


def main():
    with open(PROGRESS_PATH) as f:
        progress = json.load(f)

    with open(INPUT_CSV) as f:
        reader = csv.DictReader(f)
        original_fields = list(reader.fieldnames)
        rows = list(reader)

    # Insert email columns right after the Listing/Buyer Broker Phone columns
    new_fields = []
    for field in original_fields:
        new_fields.append(field)
        if field == 'Listing Broker Phone':
            new_fields.append('Listing Broker Email')
        elif field == 'Buyers Broker Phone':
            new_fields.append('Buyers Broker Email')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Export041026'

    # Header row with formatting
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    header_align = Alignment(horizontal='left', vertical='center')
    for col_idx, field in enumerate(new_fields, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        # Look up emails
        lfn = row.get('Listing Broker Agent First Name', '').strip()
        lln = row.get('Listing Broker Agent Last Name', '').strip()
        lco = row.get('Listing Broker Company', '').strip()
        lkey = f"{lfn}|{lln}|{lco}"
        listing_email = (progress.get(lkey) or {}).get('email') or ''

        bfn = row.get('Buyers Broker Agent First Name', '').strip()
        bln = row.get('Buyers Broker Agent Last Name', '').strip()
        bco = row.get('Buyers Broker Company', '').strip()
        bkey = f"{bfn}|{bln}|{bco}"
        buyer_email = (progress.get(bkey) or {}).get('email') or ''

        for col_idx, field in enumerate(new_fields, start=1):
            if field == 'Listing Broker Email':
                value = listing_email
            elif field == 'Buyers Broker Email':
                value = buyer_email
            else:
                value = row.get(field, '')
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Auto-width based on content (capped)
    for col_idx, field in enumerate(new_fields, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(field)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2, max_row=min(len(rows)+1, 200)):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 45)

    wb.save(OUTPUT_XLSX)

    total_with = sum(1 for v in progress.values() if v.get('email'))
    print(f"Wrote {len(rows)} rows × {len(new_fields)} columns to:")
    print(f"  {OUTPUT_XLSX}")
    print(f"Unique brokers with email: {total_with}/{len(progress)}")


if __name__ == '__main__':
    main()
